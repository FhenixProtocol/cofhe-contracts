#!/usr/bin/env python3
"""Write results/gas-measurements.xlsx (for Google Sheets) from the same data and checks as render.py.

Needs openpyxl. Run from contracts/internal/gas-measurements after render.py passes.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import render

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FONT = Font(bold=True, size=12)
GAS = "#,##0"


def header(ws, row, names):
    for col, name in enumerate(names, 1):
        c = ws.cell(row=row, column=col, value=name)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, values):
    for col, w in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def rounded(gas):
    return int(round(gas / 100.0)) * 100


def per_op_sheet(wb, rows, ops):
    ws = wb.active
    ws.title = "Gas per op"
    header(ws, 1, ["Category", "Operation", "first (gas)", "extra (gas)", "Types", "Type spread (gas)"])
    grouped = render.groups(ops)
    r = 2
    for title, entries in render.TABLES:
        for key, label, types in entries:
            ids = grouped[key]
            values = [title, label.replace("`", ""), rounded(max(rows[i]["first"] for i in ids)),
                      rounded(max(rows[i]["extra"] for i in ids)), types, render.spread(rows, ids, "first")]
            for col, v in enumerate(values, 1):
                c = ws.cell(row=r, column=col, value=v)
                if col in (3, 4, 6):
                    c.number_format = GAS
            r += 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:F{r - 1}"
    widths(ws, [30, 48, 13, 13, 34, 17])


def token_sheet(wb, tokens):
    ws = wb.create_sheet("Token transfers", 0)
    r = 1
    for name, (title, what, token_rows) in render.TOKENS.items():
        commit = tokens[name]["sepolia"]["sourceCommit"][:7]
        ws.cell(row=r, column=1, value=title).font = SECTION_FONT
        ws.cell(row=r + 1, column=1, value=f"{what} @ fhenix-confidential-contracts {commit}".replace("`", ""))
        header(ws, r + 2, ["Action", "Gas (full transaction)"])
        rows = tokens[name]["sepolia"]["rows"]
        for i, (rid, label) in enumerate(token_rows, r + 3):
            ws.cell(row=i, column=1, value=label.replace("`", ""))
            ws.cell(row=i, column=2, value=rounded(rows[rid])).number_format = GAS
        r += len(token_rows) + 5
    ws.cell(row=r, column=1, value="Full transaction gas: 21,000 base + calldata + execution. On Arbitrum, add the "
                                   "L1 data fee for the calldata. Same values on both chains.")
    widths(ws, [80, 22])


def zama_sheet(wb, zama, fh_rows, tokens):
    ws = wb.create_sheet("Zama comparison")
    meta = json.loads(Path("zama/results/zama-ops.json").read_text())
    chains = render.ZAMA_CHAINS
    ws.cell(row=1, column=1, value="FHE operations, euint64 (first / extra gas)").font = SECTION_FONT
    cols = ["Operation", "Fhenix first", "Fhenix extra"]
    for _, label in chains:
        cols += [f"{label} first", f"{label} extra"]
    header(ws, 2, cols + ["Note"])
    r = 3
    for cat in render.ZAMA_CATEGORIES:
        for m in (m for m in meta if m["category"] == cat):
            rid = m["id"]
            fh = fh_rows.get(f"{rid}__euint64")
            values = [rid, rounded(fh["first"]) if fh else None, rounded(fh["extra"]) if fh else None]
            for key, _ in chains:
                values += [rounded(zama[key]["ops"][rid]["first"]), rounded(zama[key]["ops"][rid]["extra"])]
            values.append(render.ZAMA_NOTES.get(rid, "").replace("`", ""))
            for col, v in enumerate(values, 1):
                c = ws.cell(row=r, column=col, value=v)
                if isinstance(v, int):
                    c.number_format = GAS
            r += 1
    r += 1
    ws.cell(row=r, column=1, value="Confidential token, whole transaction gas").font = SECTION_FONT
    header(ws, r + 1, ["Action", "Fhenix FHERC20"] + [f"{label} ERC7984" for _, label in chains])
    fherc20 = tokens["fherc20"]["sepolia"]["rows"]
    for i, (rid, label) in enumerate(render.ZAMA_TOKEN_ROWS, r + 2):
        values = [label.replace("`", ""), rounded(fherc20[rid])] + [rounded(zama[k]["token"][rid]) for k, _ in chains]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=i, column=col, value=v)
            if isinstance(v, int):
                c.number_format = GAS
    r = r + 3 + len(render.ZAMA_TOKEN_ROWS)
    for k, label in chains:
        z = zama[k]
        ws.cell(row=r, column=1, value=f"{label}: block {z['forkBlock']:,}, {z['executor']['version']}, "
                                       f"{z['acl']['version']}, {z['inputVerifier']['version']}, "
                                       f"{z['hcuLimit']['version']}, {z['inputSignatures']} input signature(s)")
        r += 1
    ws.cell(row=r, column=1, value="Zama: @fhevm/solidity 0.11.1, @openzeppelin/confidential-contracts 0.5.3. "
                                   "HCU limits are not gas and are not shown.")
    widths(ws, [30, 13, 13, 18, 18, 18, 18, 60])


def batch_sheet(wb, rows):
    ws = wb.create_sheet("Batches")
    ws.cell(row=1, column=1, value="Batch encrypted inputs (FHE.asEuint32s, one signature)").font = SECTION_FONT
    header(ws, 2, ["Inputs in batch"] + [int(n) for n in render.BATCH_SIZES])
    for i, (label, div) in enumerate([("Total gas", False), ("Gas per input", True)], 3):
        ws.cell(row=i, column=1, value=label)
        for col, n in enumerate(render.BATCH_SIZES, 2):
            gas = rows[f"inputBatch__{n}"]["first"] / (int(n) if div else 1)
            ws.cell(row=i, column=col, value=rounded(gas)).number_format = GAS
    ws.cell(row=6, column=1, value="Batch decryption results (euint32), total gas").font = SECTION_FONT
    header(ws, 7, ["Results in batch"] + [int(n) for n in render.BATCH_SIZES])
    for i, (op, fn) in enumerate([("publishDecryptBatch", "publishDecryptResultBatch"),
                                  ("verifyDecryptBatch", "verifyDecryptResultBatch"),
                                  ("verifyDecryptBatchSafe", "verifyDecryptResultBatchSafe")], 8):
        ws.cell(row=i, column=1, value=fn)
        for col, n in enumerate(render.BATCH_SIZES, 2):
            ws.cell(row=i, column=col, value=rounded(rows[f"{op}__{n}"]["first"])).number_format = GAS
    widths(ws, [34, 12, 12, 12, 12])


def raw_sheet(wb, chains, ops):
    ws = wb.create_sheet("Per type (raw)")
    header(ws, 1, ["Category", "Operation", "Type", "Sepolia first", "Sepolia extra",
                   "Arb Sepolia first", "Arb Sepolia extra"])
    s, a = chains["sepolia"]["rows"], chains["arbitrum-sepolia"]["rows"]
    for r, o in enumerate(ops, 2):
        values = [o["category"], o["op"], o["type"].replace("__", " → "),
                  s[o["id"]]["first"], s[o["id"]]["extra"], a[o["id"]]["first"], a[o["id"]]["extra"]]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=r, column=col, value=v)
            if col >= 4:
                c.number_format = GAS
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:G{len(ops) + 1}"
    widths(ws, [14, 24, 22, 14, 14, 17, 17])


def setup_sheet(wb, chains, deployed, exact, total):
    ws = wb.create_sheet("Setup")
    lines = [line.strip("|").split("|") for line in render.setup_table(chains, deployed).splitlines()]
    lines = [lines[0]] + lines[2:]  # drop the markdown separator row
    header(ws, 1, [c.strip() for c in lines[0]])
    for r, line in enumerate(lines[1:], 2):
        for col, v in enumerate(line, 1):
            ws.cell(row=r, column=col, value=v.strip().replace("`", "")).alignment = Alignment(wrap_text=True)
    r = len(lines) + 2
    ws.cell(row=r, column=1, value=f"Both chains run the same code; Arbitrum Sepolia: {exact} of {total} rows "
                                   "equal Sepolia to the gas unit, the rest differ by less than the rounding.")
    widths(ws, [44, 50, 50])


def notes_sheet(wb):
    ws = wb.create_sheet("Notes")
    notes = [
        "first: gas of the FHE call when it is the first FHE call in the transaction (cold TaskManager and ACL).",
        "extra: gas of the same op again in the same transaction, on operands the first call did not use "
        "(warm TaskManager, cold operand ACL entries).",
        "Gas per op: the highest value across the listed types, rounded to the nearest 100 gas. "
        "Type spread = highest − lowest raw first across those types.",
        "Arbitrum: the numbers are L2 execution gas. Arbitrum also charges an L1 data fee for the calldata of the "
        "user's transaction; encrypted inputs add calldata, FHE ops inside a contract do not.",
        "",
    ] + render.FOOTNOTES.strip().splitlines()
    for r, text in enumerate(notes, 1):
        ws.cell(row=r, column=1, value=text).alignment = Alignment(wrap_text=True)
    widths(ws, [140])


def main():
    ops = json.loads(Path("results/ops.json").read_text())
    chains = {key: json.loads(Path(f"results/{key}.json").read_text()) for key, _ in render.CHAINS}
    deployed = json.loads(Path("results/deployed.json").read_text())
    render.check_consistency(chains, deployed)
    exact, total = render.check_chains_match(chains)
    rows = chains["sepolia"]["rows"]
    render.check_type_spread(rows, ops)
    tokens = render.load_tokens()
    render.check_tokens(tokens, deployed)

    wb = Workbook()
    per_op_sheet(wb, rows, ops)
    batch_sheet(wb, rows)
    token_sheet(wb, tokens)
    zama_sheet(wb, render.load_zama(), rows, tokens)
    raw_sheet(wb, chains, ops)
    setup_sheet(wb, chains, deployed, exact, total)
    notes_sheet(wb)
    wb.save("results/gas-measurements.xlsx")
    print("wrote results/gas-measurements.xlsx")


if __name__ == "__main__":
    main()
